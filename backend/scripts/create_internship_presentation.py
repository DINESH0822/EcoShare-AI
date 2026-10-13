"""Build the editable EcoShare AI internship presentation.

Requires python-pptx in the selected Python environment. All personal details are
placeholders; project and model facts are grounded in the repository.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "EcoShare_AI_Internship_Presentation.pptx"
WIDTH, HEIGHT = 13.333, 7.5
FOREST = RGBColor(19, 77, 53)
GREEN = RGBColor(24, 125, 76)
LEAF = RGBColor(115, 176, 95)
MINT = RGBColor(230, 243, 231)
PALE = RGBColor(247, 249, 242)
WHITE = RGBColor(255, 255, 255)
INK = RGBColor(34, 54, 43)
MUTED = RGBColor(83, 105, 91)
AMBER = RGBColor(156, 90, 13)
LINE = RGBColor(211, 227, 213)


def add_text(slide, text, x, y, w, h, size=16, color=INK, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.05):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.alignment = align
        paragraph.space_after = Pt(3)
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color
    return box


def add_card(slide, x, y, w, h, fill=WHITE, line=LINE, radius=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    return shape


def add_base(prs, title, kicker, slide_no, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PALE
    rail = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.16), Inches(HEIGHT))
    rail.fill.solid()
    rail.fill.fore_color.rgb = GREEN
    rail.line.fill.background()
    add_text(slide, kicker.upper(), 0.55, 0.25, 10.5, 0.28, 10, GREEN, True)
    add_text(slide, title, 0.55, 0.55, 12.0, 0.58, 25, FOREST, True)
    if subtitle:
        add_text(slide, subtitle, 0.58, 1.12, 12.0, 0.4, 11, MUTED)
    footer = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(7.12), Inches(12.2), Inches(0.012))
    footer.fill.solid()
    footer.fill.fore_color.rgb = LINE
    footer.line.fill.background()
    add_text(slide, "EcoShare AI  |  Internship project  |  Replace editable placeholders before submission",
             0.58, 7.18, 10.8, 0.18, 8, MUTED)
    add_text(slide, f"{slide_no:02d} / 10", 11.95, 7.17, 0.75, 0.2, 8, MUTED, align=PP_ALIGN.RIGHT)
    return slide


def add_bullet_card(slide, x, y, w, h, title, body, fill=WHITE, title_size=15, body_size=12):
    add_card(slide, x, y, w, h, fill=fill)
    add_text(slide, title, x + 0.2, y + 0.16, w - 0.4, 0.34, title_size, FOREST, True)
    add_text(slide, body, x + 0.2, y + 0.58, w - 0.4, h - 0.72, body_size, INK)


def add_arrow(slide, x, y, w, h, direction="right"):
    arrow = MSO_SHAPE.RIGHT_ARROW if direction == "right" else MSO_SHAPE.DOWN_ARROW
    shape = slide.shapes.add_shape(arrow, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = LEAF
    shape.line.fill.background()


def build():
    prs = Presentation()
    prs.slide_width = Inches(WIDTH)
    prs.slide_height = Inches(HEIGHT)
    prs.core_properties.title = "EcoShare AI Internship Presentation"
    prs.core_properties.subject = "Editable project presentation based on verified repository evidence"
    prs.core_properties.author = "EcoShare AI project team"

    # 1. Cover and editable identity fields.
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PALE
    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(4.15), Inches(HEIGHT))
    panel.fill.solid()
    panel.fill.fore_color.rgb = FOREST
    panel.line.fill.background()
    add_text(slide, "EcoShare", 0.55, 0.6, 3.0, 0.52, 30, WHITE, True)
    add_text(slide, "AI", 0.58, 1.14, 1.4, 0.6, 32, RGBColor(174, 222, 147), True)
    add_text(slide, "Good food.\nBetter futures.", 0.58, 2.25, 3.0, 1.25, 25, WHITE, True)
    add_text(slide, "A role-based food rescue coordination platform", 0.6, 3.68, 3.05, 0.75, 15, RGBColor(225, 241, 226))
    add_text(slide, "INTERNSHIP PROJECT PRESENTATION", 4.75, 0.72, 7.8, 0.34, 11, GREEN, True)
    add_text(slide, "EcoShare AI", 4.75, 1.16, 7.8, 0.72, 34, FOREST, True)
    add_text(slide, "AI-Powered Food Rescue Network", 4.78, 1.98, 7.7, 0.42, 19, MUTED)
    add_card(slide, 4.8, 2.7, 3.1, 2.45, fill=WHITE)
    add_text(slide, "EDITABLE DETAILS", 5.0, 2.92, 2.7, 0.3, 10, GREEN, True)
    add_text(slide, "Student: [Your name]\nCollege: [College name]\nGuide: [Guide name]\nInternship dates: [Add verified dates]",
             5.0, 3.35, 2.72, 1.52, 14, INK)
    add_card(slide, 8.35, 2.7, 3.75, 2.45, fill=MINT)
    add_text(slide, "LOGO PLACEHOLDERS", 8.6, 2.98, 3.2, 0.3, 10, GREEN, True)
    add_text(slide, "[Insert college logo]\n\n[Insert approved project logo]",
             8.62, 3.45, 3.1, 1.25, 15, FOREST, True, align=PP_ALIGN.CENTER)
    add_text(slide, "Replace bracketed fields; no personal or college details were assumed.",
             4.8, 5.65, 7.4, 0.45, 11, MUTED)
    add_text(slide, "01 / 10", 11.95, 7.17, 0.75, 0.2, 8, MUTED, align=PP_ALIGN.RIGHT)

    # 2. Abstract.
    slide = add_base(prs, "Abstract", "Project summary", 2,
                     "A coordination platform prototype with explainable matching and clearly bounded ML claims.")
    add_bullet_card(slide, 0.65, 1.75, 3.85, 3.55, "Problem",
                    "Surplus food is time-sensitive. Donors and organizations need a practical way to discover suitable food, coordinate claims, and follow pickup and delivery states.")
    add_bullet_card(slide, 4.72, 1.75, 3.85, 3.55, "Approach",
                    "EcoShare connects donor listings, verified NGO eligibility, factor-based recommendations, role-authorized claims, and rescue-operation tracking in one web application.")
    add_bullet_card(slide, 8.79, 1.75, 3.85, 3.55, "Technical contribution",
                    "A backend rules engine ranks eligible NGOs. A separate Python Random Forest estimates claim probability; its current training examples and labels are synthetic, so saved metrics are not real-world evidence.", fill=MINT)
    add_text(slide, "Scope note: the repository demonstrates implemented software flows; it does not establish measured community or environmental outcomes.",
             0.8, 5.75, 11.8, 0.62, 13, AMBER, True)

    # 3. Introduction.
    slide = add_base(prs, "Introduction & objectives", "Context", 3)
    add_bullet_card(slide, 0.7, 1.55, 5.85, 2.1, "The coordination challenge",
                    "Food has a limited rescue window. Discovery, organization suitability, contact, and status tracking can be fragmented across manual channels.")
    add_bullet_card(slide, 6.78, 1.55, 5.85, 2.1, "Project objective",
                    "Support a traceable digital workflow that helps authorized participants coordinate redistribution of surplus food.")
    add_bullet_card(slide, 0.7, 3.9, 3.85, 2.1, "Discover",
                    "Record food, quantity, location, storage details, and expiry.")
    add_bullet_card(slide, 4.72, 3.9, 3.85, 2.1, "Coordinate",
                    "Rank eligible NGOs and authorize claims by role.")
    add_bullet_card(slide, 8.74, 3.9, 3.85, 2.1, "Track",
                    "Record pickup, delivery, and history states.")
    add_text(slide, "No live impact total or achieved waste reduction is claimed in this presentation.",
             0.82, 6.35, 11.7, 0.35, 11, AMBER, True)

    # 4. Existing system and proposed system.
    slide = add_base(prs, "Existing context vs proposed system", "Problem and solution", 4)
    add_bullet_card(slide, 0.8, 1.65, 5.7, 4.45, "Existing coordination context",
                    "• Food offers and recipient needs may be communicated separately\n\n• Matching and logistics can depend on manual comparison\n\n• Pickup progress may be difficult to trace across participants\n\nBaseline caveat: this describes the project problem context, not a measured comparison against named platforms.")
    add_bullet_card(slide, 6.82, 1.65, 5.7, 4.45, "Proposed EcoShare workflow",
                    "• Shared donation and organization records\n\n• Verification-aware eligibility and explainable recommendations\n\n• JWT and role-based authorization for claims and operations\n\n• Status and history views for rescue coordination", fill=MINT)

    # 5. Architecture.
    slide = add_base(prs, "System architecture & modules", "Implementation", 5,
                     "Existing repository stack; the Python model is separate from the rules-based matching engine.")
    boxes = [
        (0.7, "React + Vite\nrole-based UI"),
        (3.25, "Express REST API\nJWT + role checks"),
        (5.8, "MongoDB\nMongoose records"),
        (8.35, "Matching rules\nrank + reasons"),
        (10.9, "Python model\nclaim estimate"),
    ]
    for x, label in boxes:
        add_card(slide, x, 2.12, 2.05, 1.15, fill=WHITE)
        add_text(slide, label, x + 0.12, 2.35, 1.81, 0.66, 13, FOREST, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    for x in (2.82, 5.37, 7.92, 10.47):
        add_arrow(slide, x, 2.52, 0.33, 0.28)
    add_bullet_card(slide, 0.8, 3.9, 3.75, 1.65, "Core modules",
                    "Donations, organizations, claims, operations, analytics, history, notifications, and EcoBot.")
    add_bullet_card(slide, 4.8, 3.9, 3.75, 1.65, "Live services",
                    "Socket.IO is integrated. External email, SMS, WhatsApp, geocoding, and route providers require configuration.")
    add_bullet_card(slide, 8.8, 3.9, 3.75, 1.65, "Persistence & access",
                    "MongoDB stores platform records; API authorization scopes operations by authenticated role.", fill=MINT)
    add_text(slide, "Maps and development routing may show client maps or estimates without proving a live route/tile service.",
             0.85, 6.05, 11.4, 0.42, 11, AMBER, True)

    # 6. Methodology / AIML.
    slide = add_base(prs, "Methodology: explainable matching + separate prediction", "AIML workflow", 6)
    steps = [
        ("1", "Donation\nrecorded"),
        ("2", "Eligibility\nfilters"),
        ("3", "Weighted\nrules score"),
        ("4", "Ranked NGOs\n+ reasons"),
        ("5", "Authorized\nclaim & rescue"),
    ]
    x_positions = [0.65, 3.17, 5.69, 8.21, 10.73]
    for (number, label), x in zip(steps, x_positions):
        add_card(slide, x, 1.7, 1.95, 1.08, fill=MINT if number == "3" else WHITE)
        add_text(slide, number, x + 0.12, 1.87, 0.3, 0.3, 13, GREEN, True)
        add_text(slide, label, x + 0.43, 1.82, 1.35, 0.72, 12, FOREST, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    for x in (2.68, 5.2, 7.72, 10.24):
        add_arrow(slide, x, 2.06, 0.34, 0.27)
    add_bullet_card(slide, 0.75, 3.25, 5.8, 2.35, "Rule-based recommendation factors",
                    "Verified/active eligibility; category and refrigeration compatibility; distance or estimated travel time; pickup radius; remaining capacity; expiry urgency; operating hours; emergency pickup; and recorded reliability.")
    add_bullet_card(slide, 6.78, 3.25, 5.8, 2.35, "Claim-probability model",
                    "RandomForestClassifier (100 trees, max depth 7). Inputs include distance, quantity, shelf life, capacity, category, urgency, and refrigeration. Synthetic formula-generated labels mean it requires real-outcome validation.", fill=MINT)
    add_text(slide, "Scores prioritize coordination; they do not guarantee a claim, delivery, or food safety. A deterministic fallback is labelled as a fallback, not an ML prediction.",
             0.85, 5.95, 11.7, 0.58, 11, AMBER, True)

    # 7. Literature survey.
    slide = add_base(prs, "Literature survey", "Verified sources, 2020-2026", 7,
                     "Selected references frame the food-waste problem and the need for interpretable AI.")
    add_bullet_card(slide, 0.72, 1.62, 3.85, 4.65, "UNEP (2024)",
                    "Food Waste Index Report 2024\n\nExpands measurement guidance and discusses approaches to food-waste reduction. It motivates accurate measurement; it does not quantify EcoShare impact.\n\nOfficial report page:\nunep.org/resources/publication/food-waste-index-report-2024", fill=WHITE, body_size=11)
    add_bullet_card(slide, 4.74, 1.62, 3.85, 4.65, "Zhu et al. (2023)",
                    "Cradle-to-grave emissions from food loss and waste represent half of total greenhouse gas emissions from food systems.\n\nNature Food, 4, 247-256.\n\nDOI: 10.1038/s43016-023-00710-3\n\nProvides lifecycle context for why food loss and waste matter.", fill=WHITE, body_size=11)
    add_bullet_card(slide, 8.76, 1.62, 3.85, 4.65, "Barredo Arrieta et al. (2020)",
                    "Explainable Artificial Intelligence (XAI): Concepts, taxonomies, opportunities and challenges toward responsible AI.\n\nInformation Fusion, 58, 82-115.\n\nDOI: 10.1016/j.inffus.2019.12.012\n\nSupports explaining model outputs and limitations.", fill=MINT, body_size=11)
    add_text(slide, "Bibliographic metadata checked against UNEP and Crossref/publisher DOI records. No source is presented as a validation study of this application.",
             0.8, 6.42, 11.8, 0.36, 10, MUTED)

    # 8. Implementation and genuine results.
    slide = add_base(prs, "Implementation & results", "Repository evidence", 8,
                     "Values below come from the checked-in model artifact and training script, not production traffic.")
    add_bullet_card(slide, 0.75, 1.55, 5.85, 2.0, "Implementation present",
                    "React/Vite UI; Express/Mongoose API; JWT role checks; explainable NGO scoring; Python Random Forest inference; and rescue status/history workflows.")
    add_bullet_card(slide, 6.75, 1.55, 5.85, 2.0, "Stored artifact holdout metrics",
                    "ROC-AUC 0.8501  |  Accuracy 0.8117\nPrecision 0.8253  |  Recall 0.9197\nF1-score 0.8700", fill=MINT)
    add_bullet_card(slide, 0.75, 3.8, 5.85, 2.1, "Required interpretation",
                    "Training script generates 3,000 synthetic examples and labels them by sampling from a hand-written logistic formula. The reported random holdout metrics measure this synthetic setup only; they are not real-world claim-success performance.")
    add_bullet_card(slide, 6.75, 3.8, 5.85, 2.1, "Evidence still needed",
                    "A governed real claim/delivery dataset, time- or organization-separated evaluation, calibration and fairness checks, partner workflow trials, and verified deployment/provider checks.", fill=WHITE)
    add_text(slide, "No confusion matrix is included in the current saved artifact. Re-run the training pipeline to regenerate reproducible synthetic metrics.",
             0.85, 6.28, 11.7, 0.45, 10, AMBER, True)

    # 9. References.
    slide = add_base(prs, "References", "Sources", 9)
    refs = [
        ("1", "United Nations Environment Programme (2024). Food Waste Index Report 2024.",
         "Official report page: https://www.unep.org/resources/publication/food-waste-index-report-2024"),
        ("2", "Zhu, J., Luo, Z., Sun, T., Li, W., Zhou, W., Wang, X., Fei, X., Tong, H., & Yin, K. (2023). Cradle-to-grave emissions from food loss and waste represent half of total greenhouse gas emissions from food systems. Nature Food, 4, 247-256.",
         "https://doi.org/10.1038/s43016-023-00710-3"),
        ("3", "Barredo Arrieta, A. et al. (2020). Explainable Artificial Intelligence (XAI): Concepts, taxonomies, opportunities and challenges toward responsible AI. Information Fusion, 58, 82-115.",
         "https://doi.org/10.1016/j.inffus.2019.12.012"),
    ]
    y = 1.55
    for number, citation, link in refs:
        add_card(slide, 0.8, y, 11.75, 1.45, fill=WHITE)
        add_text(slide, number, 1.02, y + 0.18, 0.38, 0.35, 14, GREEN, True)
        add_text(slide, citation, 1.5, y + 0.16, 10.75, 0.76, 12, INK, True)
        add_text(slide, link, 1.5, y + 0.98, 10.75, 0.3, 10, GREEN)
        y += 1.68
    add_text(slide, "Metadata cross-checked with official UNEP and DOI/Crossref records. Accessed for this presentation in 2026.",
             0.88, 6.82, 11.5, 0.22, 9, MUTED)

    # 10. Conclusion and future scope.
    slide = add_base(prs, "Conclusion & future scope", "Takeaways", 10)
    add_bullet_card(slide, 0.78, 1.55, 5.75, 2.25, "Conclusion",
                    "EcoShare provides a role-based workflow for food listings, explainable NGO recommendations, authorized claims, and rescue progress tracking. The rules engine and claim model are distinct components.")
    add_bullet_card(slide, 6.75, 1.55, 5.75, 2.25, "Next steps",
                    "Collect governed operational outcomes; retrain and evaluate outside the synthetic setup; validate complete role-based rescue flows with partners; and monitor configured services and routing.", fill=MINT)
    add_bullet_card(slide, 0.78, 4.05, 5.75, 1.72, "AIML concepts demonstrated",
                    "Supervised classification, feature encoding, train/test split, probability estimates, precision/recall/F1/ROC-AUC, interpretability, and data-provenance limits.")
    add_bullet_card(slide, 6.75, 4.05, 5.75, 1.72, "Internship progress record",
                    "A dated weekly activity log was not present in the verified project evidence. Add your actual week-by-week tasks and guide-approved dates; do not infer completed work.", fill=WHITE)
    add_text(slide, "Thank you", 5.35, 6.15, 2.6, 0.5, 20, GREEN, True, align=PP_ALIGN.CENTER)

    prs.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
